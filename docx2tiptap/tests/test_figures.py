import io
import struct
import zlib

from docx import Document
from docx2tiptap import create_docx_from_tiptap, parse_docx_with_media, to_tiptap


def png(width=2, height=2):
    def chunk(name, data):
        return struct.pack(">I", len(data)) + name + data + struct.pack(">I", zlib.crc32(name + data))
    rows = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b""))


def test_export_figure_embeds_image_in_word():
    body = {"type": "doc", "content": [{"type": "figure", "attrs": {
        "id": "f1", "assetId": "asset-1", "alt": "Two boxes", "caption": "System diagram", "width": 320,
    }}]}
    result = create_docx_from_tiptap(body, assets={"asset-1": png()})
    word = Document(io.BytesIO(result.getvalue()))
    assert len(word.inline_shapes) == 1
    assert [p.text for p in word.paragraphs if p.text] == ["System diagram"]
    assert word.inline_shapes[0]._inline.docPr.get("descr") == "Two boxes"
    elements, comments, styles, media = parse_docx_with_media(result.getvalue())
    figures = [node for node in to_tiptap(elements, comments, styles)["content"] if node["type"] == "figure"]
    assert len(figures) == 1
    assert figures[0]["attrs"]["caption"] == "System diagram"
    assert media[figures[0]["attrs"]["assetId"]]["bytes"] == png()


def test_import_image_preserves_order_and_media():
    word = Document()
    paragraph = word.add_paragraph()
    paragraph.add_run("Before")
    paragraph.add_run().add_picture(io.BytesIO(png()))
    paragraph.add_run("After")
    source = io.BytesIO()
    word.save(source)

    elements, comments, styles, media = parse_docx_with_media(source.getvalue())
    content = [node for node in to_tiptap(elements, comments, styles)["content"] if node["type"] != "rawStylesStorage"]
    assert [node["type"] for node in content] == ["paragraph", "figure", "paragraph"]
    assert content[0]["content"][0]["text"] == "Before"
    assert content[2]["content"][0]["text"] == "After"
    asset_id = content[1]["attrs"]["assetId"]
    assert media[asset_id]["bytes"] == png()


def test_import_image_inside_table_cell():
    word = Document()
    cell = word.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(png()))
    source = io.BytesIO()
    word.save(source)
    elements, comments, styles, media = parse_docx_with_media(source.getvalue())
    table = next(node for node in to_tiptap(elements, comments, styles)["content"] if node["type"] == "table")
    figure = table["content"][0]["content"][0]["content"][0]
    assert figure["type"] == "figure"
    assert media[figure["attrs"]["assetId"]]["bytes"] == png()


def test_repeated_media_uses_one_asset():
    word = Document()
    for _ in range(2):
        word.add_paragraph().add_run().add_picture(io.BytesIO(png()))
    source = io.BytesIO()
    word.save(source)
    elements, comments, styles, media = parse_docx_with_media(source.getvalue())
    figures = [node for node in to_tiptap(elements, comments, styles)["content"] if node["type"] == "figure"]
    assert len(figures) == 2
    assert figures[0]["attrs"]["assetId"] == figures[1]["attrs"]["assetId"]
    assert len(media) == 1
